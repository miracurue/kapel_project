#!/usr/bin/env python3
"""Выбирает из выгрузки ЕСКЛП препараты группы АТХ N03, N04A (корректоры), N05, N06, N07B (зависимости) и пишет config/defaults/esklp_psych_catalog.csv.

Одна строка = одно сочетание МНН + торговое наименование + форма + дозировка + номер РУ.
Использование: python3 -I scripts/esklp-extract.py <папка выгрузки> <дата ГГГГ-ММ-ДД>
"""
import csv, glob, os, re, sys, warnings
import openpyxl

warnings.simplefilter('ignore')
src, date = sys.argv[1].rstrip('/') + '/', sys.argv[2]
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'config', 'defaults', 'esklp_psych_catalog.csv')
stamp = date.replace('-', '')

nodes = {}
ws = openpyxl.load_workbook(f'{src}esklp_smnn_{stamp}.xlsx', read_only=True)[f'esklp_smnn_{stamp}']
for i, r in enumerate(ws.iter_rows(values_only=True)):
    if i < 5 or not r[1]:
        continue
    atc = [a.strip() for a in str(r[13] or '').split(';') if a.strip()]
    if any(re.match(r'N03|N04A|N05|N06|N07B', a) for a in atc):
        nodes[r[1]] = dict(mnn=r[0], atc=atc, ftg=str(r[12] or '').replace('\n', ' '), zh=r[15], narc=r[16])

rows = {}
for f in sorted(glob.glob(f'{src}esklp_klp_{stamp}_?????.xlsx')):
    w = openpyxl.load_workbook(f, read_only=True).worksheets[1]
    for i, r in enumerate(w.iter_rows(values_only=True)):
        if i < 5 or not r[0] or r[0] not in nodes:
            continue
        if r[27] or str(r[30] or '').strip():  # закончилась запись или некорректные данные
            continue
        n = nodes[r[0]]
        key = (n['mnn'], str(r[4]).strip(), str(r[6] or '').strip(), str(r[7] or '').strip(), str(r[15] or '').strip())
        rows.setdefault(key, str(r[20] or '').strip())
        n.setdefault('seen', True)

with open(out, 'w', newline='', encoding='utf-8') as fh:
    c = csv.writer(fh)
    c.writerow(['inn', 'atc_all', 'ftg', 'zhnvlp', 'narcotic', 'trade_name', 'manufacturer', 'dosage_form', 'strength', 'reg_number', 'source_date'])
    meta = {n['mnn']: n for n in nodes.values()}
    for (mnn, trade, form, dose, reg), maker in sorted(rows.items()):
        n = meta[mnn]
        c.writerow([mnn, ';'.join(n['atc']), n['ftg'], n['zh'], n['narc'], trade, maker, form, dose, reg, date])
print('rows', len(rows), 'mnn', len({k[0] for k in rows}), 'nodes without KLP', sum(1 for n in nodes.values() if 'seen' not in n))
